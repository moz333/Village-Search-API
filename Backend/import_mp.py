from openpyxl import load_workbook
import psycopg2

conn = psycopg2.connect(
    dbname="village_db",
    user="postgres",
    password="@Kiliyam6",
    host="localhost",
    port="5432"
)

cur = conn.cursor()

file_path = r"C:\zubair\Data Analysis\capstone\all-india-villages-master-list-excel\dataset\Rdir_2011_23_MADHYA_PRADESH.xlsx"

print("Opening workbook...")

wb = load_workbook(file_path, data_only=True)
ws = wb.active

print("Reading rows...")

for row in ws.iter_rows(min_row=2, values_only=True):

    try:
        # Skip empty rows
        if not row or len(row) < 8:
            continue

        state_code = str(row[0]).zfill(2)
        state_name = str(row[1]).strip()

        district_code = str(row[2]).zfill(3)
        district_name = str(row[3]).strip()

        subdistrict_code = str(row[4]).zfill(5)
        subdistrict_name = str(row[5]).strip()

        village_code = str(row[6]).zfill(6)
        village_name = str(row[7]).strip()

        if village_name.lower() == 'nan':
            continue

        # ---------------- STATE ----------------
        cur.execute("""
            INSERT INTO state(name, code, country_id)
            VALUES (%s, %s, 1)
            ON CONFLICT (code) DO NOTHING
        """, (state_name, state_code))

        cur.execute(
            "SELECT id FROM state WHERE code=%s",
            (state_code,)
        )

        state_id = cur.fetchone()[0]

        # ---------------- DISTRICT ----------------
        cur.execute("""
            INSERT INTO district(name, code, state_id)
            VALUES (%s, %s, %s)
            ON CONFLICT (code, state_id) DO NOTHING
        """, (district_name, district_code, state_id))

        cur.execute("""
            SELECT id FROM district
            WHERE code=%s AND state_id=%s
        """, (district_code, state_id))

        district_id = cur.fetchone()[0]

        # ---------------- SUBDISTRICT ----------------
        cur.execute("""
            INSERT INTO sub_district(name, code, district_id)
            VALUES (%s, %s, %s)
            ON CONFLICT (code, district_id) DO NOTHING
        """, (subdistrict_name, subdistrict_code, district_id))

        cur.execute("""
            SELECT id FROM sub_district
            WHERE code=%s AND district_id=%s
        """, (subdistrict_code, district_id))

        subdistrict_id = cur.fetchone()[0]

        # ---------------- VILLAGE ----------------
        cur.execute("""
            INSERT INTO village(name, code, sub_district_id)
            VALUES (%s, %s, %s)
            ON CONFLICT (code) DO NOTHING
        """, (village_name, village_code, subdistrict_id))

    except Exception as e:
        print("Error:", e)

conn.commit()

print("✅ Madhya Pradesh imported successfully!")

cur.close()
conn.close()